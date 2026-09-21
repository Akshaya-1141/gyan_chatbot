import os
import base64
import uuid
from io import BytesIO
from PIL import Image
import pypdf
import docx

UPLOAD_BASE = os.path.join(os.path.dirname(__file__), "uploads")
DOCS_DIR = os.path.join(UPLOAD_BASE, "documents")
IMAGES_DIR = os.path.join(UPLOAD_BASE, "images")
AUDIO_DIR = os.path.join(UPLOAD_BASE, "audio")

for d in [DOCS_DIR, IMAGES_DIR, AUDIO_DIR]:
    os.makedirs(d, exist_ok=True)

ALLOWED_DOC_EXTS = {
    ".pdf", ".docx", ".txt", ".csv", ".tsv", ".json", ".md",
    ".py", ".js", ".html", ".css", ".xml", ".yaml", ".yml", ".sh", ".sql", ".log"
}

ALLOWED_IMG_EXTS = {
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"
}

ALLOWED_AUDIO_EXTS = {
    ".mp3", ".wav", ".m4a", ".webm", ".ogg", ".aac"
}

MAX_TEXT_CHAR_LIMIT = 50000  # Safe context window allocation


def format_file_size(num_bytes):
    for unit in ['B', 'KB', 'MB', 'GB']:
        if num_bytes < 1024.0:
            return f"{num_bytes:.1f} {unit}"
        num_bytes /= 1024.0
    return f"{num_bytes:.1f} TB"


def extract_pdf_text(filepath):
    text_parts = []
    try:
        reader = pypdf.PdfReader(filepath)
        total_pages = len(reader.pages)
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"--- Page {i+1} of {total_pages} ---\n{page_text}")
        full_text = "\n\n".join(text_parts)
        if len(full_text) > MAX_TEXT_CHAR_LIMIT:
            full_text = full_text[:MAX_TEXT_CHAR_LIMIT] + "\n\n...[Content truncated due to length limits]..."
        return full_text if full_text.strip() else "[Empty or scanned image-only PDF. No extractable text found.]"
    except Exception as e:
        return f"[Error extracting PDF: {str(e)}]"


def extract_docx_text(filepath):
    try:
        doc = docx.Document(filepath)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                full_text.append(" | ".join(row_text))
        result = "\n".join(full_text)
        if len(result) > MAX_TEXT_CHAR_LIMIT:
            result = result[:MAX_TEXT_CHAR_LIMIT] + "\n\n...[Content truncated due to length limits]..."
        return result if result.strip() else "[Empty DOCX document]"
    except Exception as e:
        return f"[Error extracting DOCX: {str(e)}]"


def extract_plain_text(filepath):
    encodings = ['utf-8', 'latin-1', 'cp1252', 'utf-16']
    for enc in encodings:
        try:
            with open(filepath, 'r', encoding=enc) as f:
                content = f.read()
                if len(content) > MAX_TEXT_CHAR_LIMIT:
                    content = content[:MAX_TEXT_CHAR_LIMIT] + "\n\n...[Content truncated due to length limits]..."
                return content
        except UnicodeDecodeError:
            continue
        except Exception as e:
            return f"[Error reading text file: {str(e)}]"
    return "[Unable to decode text file with standard encodings]"


def process_uploaded_file(file_storage):
    """
    Saves and processes an uploaded file.
    Returns metadata dict with extracted content or base64 URI.
    """
    orig_filename = file_storage.filename or "uploaded_file"
    ext = os.path.splitext(orig_filename)[1].lower()
    file_id = str(uuid.uuid4())
    safe_name = f"{file_id[:8]}_{orig_filename}"

    file_bytes = file_storage.read()
    file_size = len(file_bytes)
    formatted_size = format_file_size(file_size)

    # 1. Process Images
    if ext in ALLOWED_IMG_EXTS:
        save_path = os.path.join(IMAGES_DIR, safe_name)
        with open(save_path, "wb") as f:
            f.write(file_bytes)
        
        # Validate and optimize image
        try:
            with Image.open(save_path) as img:
                img_format = img.format or "PNG"
                if img.mode in ("RGBA", "P") and ext in [".jpg", ".jpeg"]:
                    img = img.convert("RGB")
                
                # Resize if excessively large for Vision API
                max_dim = 1920
                if img.width > max_dim or img.height > max_dim:
                    img.thumbnail((max_dim, max_dim))
                    buffer = BytesIO()
                    img.save(buffer, format=img_format)
                    encoded = base64.b64encode(buffer.getvalue()).decode('utf-8')
                else:
                    encoded = base64.b64encode(file_bytes).decode('utf-8')

                mime = f"image/{ext.replace('.', '')}"
                if ext == ".jpg":
                    mime = "image/jpeg"
                data_uri = f"data:{mime};base64,{encoded}"

                return {
                    "id": file_id,
                    "filename": orig_filename,
                    "type": "image",
                    "extension": ext,
                    "size_bytes": file_size,
                    "size_formatted": formatted_size,
                    "file_path": save_path,
                    "data_uri": data_uri,
                    "text_content": ""
                }
        except Exception as e:
            return {
                "error": f"Invalid image file: {str(e)}"
            }

    # 2. Process Documents / Code files
    elif ext in ALLOWED_DOC_EXTS:
        save_path = os.path.join(DOCS_DIR, safe_name)
        with open(save_path, "wb") as f:
            f.write(file_bytes)

        if ext == ".pdf":
            extracted_text = extract_pdf_text(save_path)
        elif ext == ".docx":
            extracted_text = extract_docx_text(save_path)
        else:
            extracted_text = extract_plain_text(save_path)

        return {
            "id": file_id,
            "filename": orig_filename,
            "type": "document",
            "extension": ext,
            "size_bytes": file_size,
            "size_formatted": formatted_size,
            "file_path": save_path,
            "text_content": extracted_text,
            "data_uri": ""
        }

    # 3. Audio files
    elif ext in ALLOWED_AUDIO_EXTS:
        save_path = os.path.join(AUDIO_DIR, safe_name)
        with open(save_path, "wb") as f:
            f.write(file_bytes)

        return {
            "id": file_id,
            "filename": orig_filename,
            "type": "audio",
            "extension": ext,
            "size_bytes": file_size,
            "size_formatted": formatted_size,
            "file_path": save_path,
            "text_content": "",
            "data_uri": ""
        }

    else:
        return {
            "error": f"Unsupported file extension '{ext}'. Supported formats: PDF, DOCX, TXT, CSV, MD, Code files, Images (PNG, JPG, WEBP), Audio."
        }
